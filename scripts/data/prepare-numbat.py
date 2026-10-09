"""Offline, pinned official NUMBAT station entry/exit normalisation. No provider requests."""
import hashlib
import json
import math
import pathlib
import sys

import openpyxl

root = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "supabase/.temp/phase8-implementation").resolve()
if not root.is_relative_to(pathlib.Path("supabase/.temp").resolve()):
    raise ValueError("Ignored development directory required")
hashes = {
    "MON": "89ebca13a5f611e255c3cc31eb34917a7262c642135e943ee6714df17a108b05",
    "TWT": "a10e5f767750538ecac99cf0d2131243672574df88b9c933f03b9357c5f67db2",
    "FRI": "22eeb8fe2fd5ee2c974aaff81c7f3e114c53e39f46cd03c273c072974f5d9b5c",
    "SAT": "64cdfa78f312332740f9a562f81f5dac12114bb439d4a219bcf095ee28ad3821",
    "SUN": "c064331e50cf92edd022ef9f1f7efb9c21ca38edabb963a7bf5758682d60bc8c",
}
periods = []
for index in range(96):
    start = (300 + index * 15) % 1440
    end = (start + 15) % 1440
    periods.append(f"{start // 60:02}{start % 60:02}-{end // 60:02}{end % 60:02}")
prefix = ["NLC", "ASC", "Station", "Fare Zone", "Total", "Early", "AM Peak", "Midday", "PM Peak", "Evening", "Late"]
all_rows = []
receipts = []
for day, expected in hashes.items():
    path = root / f"NBT25{day}_Outputs.xlsx"
    if hashlib.sha256(path.read_bytes()).hexdigest() != expected:
        raise ValueError("Pinned NUMBAT workbook required")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    cover = list(workbook["_Cover"].iter_rows(values_only=True))
    # Reference year/day are pinned by artifact identity; workbook production date is separate.
    production_date = str(cover[4][1])
    if production_date != "2026-07-01 00:00:00":
        raise ValueError("NUMBAT production date changed")
    seen = set()
    summaries = {}
    for sheet, measure in [("Station_Entries", "gateline_entries"), ("Station_Exits", "gateline_exits")]:
        rows = list(workbook[sheet].iter_rows(values_only=True))
        if list(rows[2]) != prefix + periods or list(rows[1][11:]) != list(range(21, 117)):
            raise ValueError("NUMBAT native traffic-day columns changed")
        admitted = 0
        for row in rows[3:]:
            if all(value is None for value in row):
                continue
            nlc, asc, station, zone = row[:4]
            if not isinstance(nlc, int) or not isinstance(asc, str) or not isinstance(station, str):
                raise ValueError("NUMBAT station identity missing")
            key = (asc, measure)
            if key in seen:
                raise ValueError("Duplicate native station/measure")
            seen.add(key)
            values = list(row[11:])
            original_totals = list(row[4:11])
            if any(value is not None and (not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0) for value in values + original_totals):
                raise ValueError("NUMBAT invalid modelled value")
            # Retain fractions, zeros and blanks; no rounded fabricated pedestrian count.
            known = sum(value for value in values if value is not None)
            total = row[4]
            if all(value is not None for value in values) and total is not None and abs(known - total) > max(0.00001, total * 0.000001):
                raise ValueError("Native quarter-hour sum disagrees with published total")
            all_rows.append({"schemaVersion": 1, "year": 2025, "dayType": day, "measure": measure,
                "station": {"nlc": nlc, "asc": asc, "name": station, "fareZone": str(zone) if zone is not None else None},
                "trafficDayStartMinutes": 300, "quarterHourIds": list(range(21, 117)), "periodLabels": periods,
                "values": values, "missingReasons": ["source_blank" if value is None else None for value in values],
                "publishedTotals": original_totals, "workbookProducedAt": production_date,
                "sourceSheet": sheet, "sourceRow": admitted + 4,
                "sourceKind": "modelled", "units": "typical_day_gateline_passenger_movements"})
            admitted += 1
        summaries[measure] = admitted
    receipts.append({"dayType": day, "archiveSha256": expected, "rows": summaries})
    workbook.close()
all_rows.sort(key=lambda r: (r["dayType"], r["station"]["asc"], r["measure"]))
(root / "numbat-normalised.json").write_text(json.dumps(all_rows, separators=(",", ":")), encoding="utf8")
print(json.dumps({"profiles": len(all_rows), "artifacts": receipts, "activated": False}))
