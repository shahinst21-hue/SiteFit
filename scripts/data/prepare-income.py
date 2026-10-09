"""Offline pinned ONS workbook preparation. Requires read-only openpyxl;
no provider/account requests, formulas, source edits or database activation."""
import hashlib
import json
import pathlib
import sys
import openpyxl

root = pathlib.Path("supabase/.temp").resolve()
directory = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else root / "phase8-implementation").resolve()
if not directory.is_relative_to(root):
    raise RuntimeError("Ignored development artifact directory required")
archive = (directory / "income-fye2023.xlsx").read_bytes()
if len(archive) != 1863497 or hashlib.sha256(archive).hexdigest() != "9f17da5c218d2edd336106e1f219e0f7f25e82373afa8be5d0dc7622e18c7ab7":
    raise RuntimeError("Pinned FYE2023 workbook required")
lookup = (directory / "london-native-lookup.json").read_bytes()
if hashlib.sha256(lookup).hexdigest() != "54274c06548e400654adc0a438cb3bd48165dc12135963f5aa5594739ec5baa4":
    raise RuntimeError("Pinned native 2021 London lookup required")
codes = {r["MSOA21CD"] for r in json.loads(lookup)}
book = openpyxl.load_workbook(directory / "income-fye2023.xlsx", read_only=True, data_only=True)
try:
    sheet = book["Net income after housing costs"]
    header = next(sheet.iter_rows(min_row=4, max_row=4, values_only=True))
    if header != ("MSOA code", "MSOA name", "Local authority code", "Local authority name", "Region code", "Region name",
                  "Disposable (net) annual income after housing costs (£)", "Upper confidence limit (£)",
                  "Lower confidence limit (£)", "Confidence interval (£)"):
        raise RuntimeError("Income semantics changed")
    result = []
    for r in sheet.iter_rows(min_row=5, values_only=True):
        if r[4] != "E12000007":
            continue
        if r[0] not in codes or any(type(v) not in (int, float) for v in r[6:10]) or not 0 <= r[8] <= r[6] <= r[7] or r[7] - r[8] != r[9]:
            raise RuntimeError("Invalid income geography/interval")
        result.append({"code": r[0], "mean": r[6], "upper": r[7], "lower": r[8], "intervalWidth": r[9], "ladCode": r[2]})
    if len(result) != 1002 or {r["code"] for r in result} != codes:
        raise RuntimeError("Incomplete or duplicate London MSOAs")
    result.sort(key=lambda r: r["code"])
    (directory / "income-london-normalised.json").write_text(json.dumps(result, separators=(",", ":")), encoding="utf-8")
    print(json.dumps({"dataset": "income-AHC-FYE2023", "rows": len(result), "sourceArchiveVerified": True, "confidenceBoundsPreserved": True}))
finally:
    book.close()
